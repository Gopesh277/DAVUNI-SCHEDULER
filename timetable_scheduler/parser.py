"""
parser.py

Parses the department's teaching-load register (.xlsx or .csv) into a
flat list of normalized course-offering records.

The register has merged cells: S.No / Faculty Name / Employee ID /
Nature of Appointment only appear on a faculty member's first row, with
every subsequent course row for that faculty left blank in those columns.
This module forward-fills those columns before building the records that
the rest of the application consumes.

Expected column order:

    S.No | Faculty Name | Employee ID | Nature of Appointment |
    Course Code | Course Name | Course Type | Programme/s | Semester |
    Contact Hours(Theory, Tutorial, Practical) |
    Teaching Load(...) | Signature

The Teaching Load columns are informational totals in the source sheet
and are recomputed by this application rather than trusted verbatim.

Group handling
--------------

The parser recognizes group markers at the end of programme text:

    B.Tech CSE B(G1)
    B.Tech CSE B (G1)
    B.Tech CSE B-G1
    B.Tech CSE B - G1
    B.Tech CSE B:G1
    B.Tech CSE B, G1
    B.Tech CSE BG1

All of the above resolve to:

    base_programme = "B.Tech CSE B"
    group = "G1"

This is important because the timetable solver needs to understand that
G1 is part of the same parent class as the shared lecture.

Multiple groups such as:

    B.Tech CSE B(G1, G2)

are treated as a shared/joint session:

    group = None
    base_programme = "B.Tech CSE B"

"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any


# ---------------------------------------------------------------------------
# Group marker parsing
# ---------------------------------------------------------------------------
#
# Matches a trailing group marker with optional separator and optional
# parentheses/brackets.
#
# Examples:
#
#   "B.Tech CSE B(G1)"
#   "B.Tech CSE B (G1)"
#   "B.Tech CSE B-G1"
#   "B.Tech CSE B - G1"
#   "B.Tech CSE B:G1"
#   "B.Tech CSE B, G1"
#   "B.Tech CSE BG1"
#
# It also supports multiple groups:
#
#   "B.Tech CSE B(G1, G2)"
#
# The separator is deliberately included in the match so that when the
# group marker is stripped, the leftover "-", ":", "," etc. do not remain
# attached to the base programme.
#
_GROUP_RE = re.compile(
    r"""
    # Optional separator before the group marker.
    #
    # Examples:
    #   " - G1"
    #   "-G1"
    #   ": G1"
    #   ", G1"
    #
    (?:\s*[-:,\u2013\u2014]\s*)?

    # Optional opening bracket/parenthesis.
    (?:[\(\[\{]\s*)?

    # One or more group tokens.
    #
    # Examples:
    #   G1
    #   G1, G2
    #   G1 G2
    #   G1, G2, G3
    #
    (
        (?:g\s*\d+\s*,?\s*)+
    )

    # Optional closing bracket/parenthesis.
    (?:\s*[\)\]\}])?

    # Group marker must be at the end of the programme string.
    \s*$
    """,
    re.IGNORECASE | re.VERBOSE,
)


# ---------------------------------------------------------------------------
# Programme normalization
# ---------------------------------------------------------------------------

def _normalize_programme_text(raw: str) -> str:
    """
    Normalize programme text so minor formatting differences do not create
    separate sections.

    Examples:

        "B.tech CSE A"
        "B.Tech CSE A"

    become:

        "B.Tech CSE A"

    Also normalizes whitespace around '&':

        "CS &AI"
        "CS& AI"
        "CS&AI"
        "CS & AI"

    all become:

        "CS & AI"
    """

    s = (raw or "Unassigned").strip()

    # Collapse repeated whitespace.
    s = re.sub(r"\s+", " ", s)

    # Normalize whitespace around ampersand.
    s = re.sub(r"\s*&\s*", " & ", s)

    # Normalize B.Tech.
    s = re.sub(
        r"(?i)^b\.\s*tech\b",
        "B.Tech",
        s,
    )

    # Normalize M.Tech.
    s = re.sub(
        r"(?i)^m\.\s*tech\b",
        "M.Tech",
        s,
    )

    return s.strip()


# ---------------------------------------------------------------------------
# CourseOffering
# ---------------------------------------------------------------------------

@dataclass
class CourseOffering:
    """
    One faculty/course/programme offering from the teaching-load register.
    """

    faculty: str
    nature: str
    course_code: str
    course_name: str
    course_type: str
    programme: str
    semester: str

    theory_hrs: int = 0
    tutorial_hrs: int = 0
    practical_hrs: int = 0

    emp_id: Any = None

    # -----------------------------------------------------------------------
    # Group
    # -----------------------------------------------------------------------

    @property
    def group(self) -> str | None:
        """
        Return the sub-group for this course row.

        Examples:

            "B.Tech CSE B(G1)"  -> "G1"
            "B.Tech CSE B-G1"   -> "G1"
            "B.Tech CSE B - G1" -> "G1"
            "B.Tech CSE B:G1"   -> "G1"

        A row containing multiple groups, for example:

            "B.Tech CSE B(G1, G2)"

        is treated as a shared/joint class and returns None.
        """

        programme = (self.programme or "").strip()

        match = _GROUP_RE.search(programme)

        if not match:
            return None

        # Extract individual group tokens.
        #
        # Example:
        #   "G1, G2" -> ["G1", "G2"]
        #
        tokens = re.findall(
            r"g\s*\d+",
            match.group(1),
            re.IGNORECASE,
        )

        # Normalize:
        #
        #   "g 1" -> "G1"
        #   "G 2" -> "G2"
        #
        normalized_tokens = sorted(
            {
                token.upper().replace(" ", "")
                for token in tokens
            }
        )

        # No valid group found.
        if not normalized_tokens:
            return None

        # One group = actual subgroup.
        if len(normalized_tokens) == 1:
            return normalized_tokens[0]

        # Multiple groups = shared/joint session.
        return None

    # -----------------------------------------------------------------------
    # Base programme
    # -----------------------------------------------------------------------

    @property
    def base_programme(self) -> str:
        """
        Return the programme without its trailing group marker.

        Examples:

            "B.Tech CSE B(G1)"
                -> "B.Tech CSE B"

            "B.Tech CSE B - G1"
                -> "B.Tech CSE B"

            "B.Tech CSE B:G1"
                -> "B.Tech CSE B"

            "B.Tech CSE B, G1"
                -> "B.Tech CSE B"

            "B.Tech CSE B(G1, G2)"
                -> "B.Tech CSE B"
        """

        prog = (self.programme or "Unassigned").strip()

        match = _GROUP_RE.search(prog)

        if match:
            # Because _GROUP_RE includes the separator before the group,
            # match.start() points before the separator.
            prog = prog[:match.start()]

        return _normalize_programme_text(prog)

    # -----------------------------------------------------------------------
    # Base section
    # -----------------------------------------------------------------------

    @property
    def base_section(self) -> str:
        """
        Return the parent class/section independent of subgroup.

        This is the identifier used by the timetable solver to detect
        conflicts between shared lectures and group-specific practicals.

        Example:

            B.Tech CSE B(G1), semester 5
                -> B.Tech CSE B - 5 Sem

            B.Tech CSE B, semester 5
                -> B.Tech CSE B - 5 Sem

        Therefore the solver understands that G1 belongs to CSE B.
        """

        sem = (self.semester or "").strip()

        if sem:
            return f"{self.base_programme} - {sem} Sem"

        return self.base_programme

    # -----------------------------------------------------------------------
    # Exact section
    # -----------------------------------------------------------------------

    @property
    def section(self) -> str:
        """
        Return the exact schedulable unit.

        A normal class:

            B.Tech CSE B - 5 Sem

        A group-specific class:

            B.Tech CSE B (G1) - 5 Sem

        This distinction allows:

            base_section
                -> conflict parent

            section
                -> actual schedulable unit
        """

        sem = (self.semester or "").strip()

        if self.group:
            label = f"{self.base_programme} ({self.group})"
        else:
            label = self.base_programme

        if sem:
            return f"{label} - {sem} Sem"

        return label

    # -----------------------------------------------------------------------
    # Dictionary representation
    # -----------------------------------------------------------------------

    def to_dict(self) -> dict:
        """
        Convert this CourseOffering into a dictionary.
        """

        return asdict(self)


# ---------------------------------------------------------------------------
# Utility conversion functions
# ---------------------------------------------------------------------------

def _to_int(v) -> int:
    """
    Safely convert a spreadsheet value into an integer.

    Examples:

        3       -> 3
        3.0     -> 3
        "3"     -> 3
        "3.0"   -> 3
        ""      -> 0
        None    -> 0
        "abc"   -> 0
    """

    try:
        if v is None or v == "":
            return 0

        return int(float(v))

    except (TypeError, ValueError):
        return 0


def _clean(v) -> str | None:
    """
    Convert a cell value to stripped text.

    Empty cells become None.
    """

    if v is None:
        return None

    s = str(v).strip()

    return s if s else None


# ---------------------------------------------------------------------------
# XLSX reader
# ---------------------------------------------------------------------------

def _rows_from_xlsx(path: Path) -> list[list]:
    """
    Read an XLSX/XLSM file and return all rows as lists.

    Uses the first sheet unless a sheet named 'Sheet1' exists.
    """

    import openpyxl

    wb = openpyxl.load_workbook(
        path,
        data_only=True,
    )

    sheet_name = (
        "Sheet1"
        if "Sheet1" in wb.sheetnames
        else wb.sheetnames[0]
    )

    ws = wb[sheet_name]

    return [
        list(row)
        for row in ws.iter_rows(values_only=True)
    ]


# ---------------------------------------------------------------------------
# CSV reader
# ---------------------------------------------------------------------------

def _rows_from_csv(path: Path) -> list[list]:
    """
    Read a CSV file.

    utf-8-sig allows normal UTF-8 files as well as files containing
    a UTF-8 BOM.
    """

    with open(
        path,
        newline="",
        encoding="utf-8-sig",
    ) as f:

        return [
            row
            for row in csv.reader(f)
        ]


# ---------------------------------------------------------------------------
# Header detection
# ---------------------------------------------------------------------------

def _find_header_row(rows: list[list]) -> int:
    """
    Locate the header row by searching for the 'Course Code' column.
    """

    for i, row in enumerate(rows):

        if not row:
            continue

        for cell in row:

            if cell is None:
                continue

            if str(cell).strip() == "Course Code":
                return i

    # Fall back to the first row.
    return 0


# ---------------------------------------------------------------------------
# Main parser
# ---------------------------------------------------------------------------

def parse_teaching_load(
    path: str | Path,
) -> list[CourseOffering]:
    """
    Parse a teaching-load register into CourseOffering records.

    Supported formats:

        .xlsx
        .xlsm
        .csv

    The parser handles merged-cell style registers where faculty
    information appears only on the first row of each faculty member.

    Raises:
        ValueError:
            If the file format is unsupported or no course rows are found.
    """

    path = Path(path)

    # -----------------------------------------------------------------------
    # Read source file
    # -----------------------------------------------------------------------

    suffix = path.suffix.lower()

    if suffix in (".xlsx", ".xlsm"):
        rows = _rows_from_xlsx(path)

    elif suffix == ".csv":
        rows = _rows_from_csv(path)

    else:
        raise ValueError(
            f"Unsupported file type: {path.suffix}. "
            "Use .xlsx or .csv."
        )

    # -----------------------------------------------------------------------
    # Find header
    # -----------------------------------------------------------------------

    header_idx = _find_header_row(rows)

    # The department template normally contains:
    #
    #   header row
    #   units/sub-header row
    #   actual data
    #
    # Therefore skip both header and units row.
    #
    data_rows = rows[header_idx + 2:]

    # -----------------------------------------------------------------------
    # Output
    # -----------------------------------------------------------------------

    offerings: list[CourseOffering] = []

    # Current faculty information.
    #
    # These are forward-filled when the spreadsheet uses merged cells.
    cur_faculty = None
    cur_emp = None
    cur_nature = None

    # -----------------------------------------------------------------------
    # Process rows
    # -----------------------------------------------------------------------

    for row in data_rows:

        # Completely empty row.
        if not row:
            continue

        if all(
            c is None or str(c).strip() == ""
            for c in row
        ):
            continue

        # Make sure we always have at least 12 columns.
        padded = (
            list(row)
            + [None] * max(0, 12 - len(row))
        )

        (
            sno,
            fname,
            emp,
            nature,
            ccode,
            cname,
            ctype,
            prog,
            sem,
            th,
            tu,
            pr,
        ) = padded[:12]

        # -------------------------------------------------------------------
        # Forward-fill faculty information
        # -------------------------------------------------------------------
        #
        # The register may look like:
        #
        # S.No | Faculty | Emp ID | Nature | Course
        #  1   | Rahul   | 123    | Regular| CS101
        #      |         |        |        | CS102
        #      |         |        |        | CS103
        #
        # The blank faculty rows belong to Rahul.
        #

        if _clean(sno) is not None:

            cur_faculty = _clean(fname)

            cur_emp = emp

            cur_nature = (
                _clean(nature)
                or "Regular"
            )

        # -------------------------------------------------------------------
        # Determine whether this is an empty/non-course row
        # -------------------------------------------------------------------

        is_empty_course = not any(
            [
                _clean(ccode),
                _clean(cname),
                _clean(prog),
                _clean(sem),
                _to_int(th),
                _to_int(tu),
                _to_int(pr),
            ]
        )

        # Ignore rows without a faculty.
        #
        # This also protects against unrelated rows before the first
        # faculty record.
        if is_empty_course or not cur_faculty:
            continue

        # -------------------------------------------------------------------
        # Build CourseOffering
        # -------------------------------------------------------------------

        offering = CourseOffering(
            faculty=cur_faculty,

            nature=(
                cur_nature
                or "Regular"
            ),

            course_code=(
                _clean(ccode)
                or "-"
            ),

            course_name=(
                _clean(cname)
                or "Untitled course"
            ),

            course_type=(
                _clean(ctype)
                or ""
            ),

            programme=(
                _clean(prog)
                or "Unassigned"
            ),

            semester=(
                _clean(sem)
                or ""
            ),

            theory_hrs=_to_int(th),

            tutorial_hrs=_to_int(tu),

            practical_hrs=_to_int(pr),

            emp_id=cur_emp,
        )

        offerings.append(offering)

    # -----------------------------------------------------------------------
    # No results
    # -----------------------------------------------------------------------

    if not offerings:

        raise ValueError(
            f"No course rows found in {path}. "
            "Expected the standard department register columns "
            "(Faculty Name ... Course Code ... Contact Hours)."
        )

    return offerings


# ---------------------------------------------------------------------------
# Bytes parser
# ---------------------------------------------------------------------------

def parse_teaching_load_bytes(
    data: bytes,
    filename: str,
) -> list[CourseOffering]:
    """
    Parse an uploaded teaching-load register directly from bytes.

    Useful for FastAPI UploadFile.

    The data is temporarily written to disk so that the normal
    parse_teaching_load() implementation can be reused.
    """

    import tempfile

    suffix = (
        Path(filename).suffix
        or ".xlsx"
    )

    with tempfile.NamedTemporaryFile(
        suffix=suffix,
        delete=False,
    ) as tmp:

        tmp.write(data)

        tmp_path = tmp.name

    try:

        return parse_teaching_load(tmp_path)

    finally:

        Path(tmp_path).unlink(
            missing_ok=True
        )
